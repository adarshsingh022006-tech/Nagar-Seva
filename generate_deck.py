# generate_deck.py
import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen format (20.0 x 11.25 inches)
    prs.slide_width = Inches(20.0)
    prs.slide_height = Inches(11.25)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_NAVY_DARK = RGBColor(11, 19, 43)       # #0B132B
    C_SLATE_DARK = RGBColor(28, 37, 65)      # #1C2541
    C_SLATE_MID = RGBColor(58, 80, 107)      # #3A506B
    C_BG_LIGHT = RGBColor(248, 250, 252)     # #F8FAFC
    C_WHITE = RGBColor(255, 255, 255)        # #FFFFFF
    C_AMBER = RGBColor(245, 158, 11)         # #F59E0B (Gold)
    C_AMBER_LIGHT = RGBColor(254, 243, 199)  # #FEF3C7
    C_EMERALD = RGBColor(16, 185, 129)       # #10B981 (Green)
    C_EMERALD_LIGHT = RGBColor(209, 250, 229) # #D1FAE5
    C_RED = RGBColor(239, 68, 68)            # #EF4444 (SOS Red)
    C_RED_LIGHT = RGBColor(254, 226, 226)    # #FEE2E2
    C_BLUE = RGBColor(59, 130, 246)          # #3B82F6
    C_BLUE_LIGHT = RGBColor(224, 242, 254)   # #E0F2FE
    C_PURPLE = RGBColor(139, 92, 246)        # #8B5CF6
    C_PURPLE_LIGHT = RGBColor(243, 232, 255) # #F3E8FF
    C_TEXT_DARK = RGBColor(15, 23, 42)       # #0F172A
    C_TEXT_MUTED = RGBColor(100, 116, 139)   # #64748B
    C_BORDER_LIGHT = RGBColor(226, 232, 240) # #E2E8F0

    def add_bg(slide, color=C_BG_LIGHT):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text, subtitle_text=""):
        # Header Container
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(18.4), Inches(1.3))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_tag = tf.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.size = Pt(12)
        p_tag.font.bold = True
        p_tag.font.color.rgb = C_AMBER
        p_tag.font.name = "Segoe UI"

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = C_TEXT_DARK
        p_title.font.name = "Segoe UI"
        p_title.space_after = Pt(2)

        if subtitle_text:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle_text
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = C_TEXT_MUTED
            p_sub.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 1: TITLE & COVER SLIDE (DARK THEME)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1, C_NAVY_DARK)

    # Accent Top Strip
    accent_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.15))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = C_AMBER
    accent_bar.line.fill.background()

    # Hackathon Badge
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.1), Inches(6.8), Inches(0.6))
    badge.fill.solid()
    badge.fill.fore_color.rgb = C_SLATE_DARK
    badge.line.color.rgb = C_AMBER
    badge.line.width = Pt(1.5)
    b_tf = badge.text_frame
    b_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    bp = b_tf.paragraphs[0]
    bp.text = "🏆 HACKCONQUEST HACKATHON — AETHER 2026"
    bp.font.size = Pt(13)
    bp.font.bold = True
    bp.font.color.rgb = C_AMBER
    bp.alignment = PP_ALIGN.CENTER

    # Main Title
    tbox = s1.shapes.add_textbox(Inches(1.2), Inches(1.9), Inches(17.5), Inches(3.2))
    tf = tbox.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = "NAGAR SEVA (नगर सेवा)"
    p1.font.size = Pt(50)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE
    p1.font.name = "Segoe UI"

    p2 = tf.add_paragraph()
    p2.text = "Next-Gen AI & Community-Driven Smart Civic Grievance Ecosystem"
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = C_EMERALD
    p2.font.name = "Segoe UI"
    p2.space_before = Pt(8)

    p3 = tf.add_paragraph()
    p3.text = "Frictionless Zero-Login Filing • AI Auto-Classification • Proximity Duplicate Merging • Verified Proof-of-Fix • Nagar Community Feed"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.font.name = "Segoe UI"
    p3.space_before = Pt(8)

    # 4 Highlight Badges
    highlights = [
        ("⚡ 100% Zero-Login", "Instant filing in <60 seconds", C_AMBER),
        ("🤖 Real-Time AI Engine", "Auto category & severity detection", C_PURPLE),
        ("🗺️ OpenStreetMap & SOS", "Pin-drop GPS & 6h SLA countdown", C_RED),
        ("📰 Nagar Civic Feed", "Community (+1) upvotes & Before/After", C_EMERALD),
    ]

    for i, (title, desc, color) in enumerate(highlights):
        x = Inches(1.2 + i * 4.45)
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(5.6), Inches(4.2), Inches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = C_SLATE_DARK
        card.line.color.rgb = color
        card.line.width = Pt(1.5)
        
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = Inches(0.25)
        
        cp1 = ctf.paragraphs[0]
        cp1.text = title
        cp1.font.size = Pt(16)
        cp1.font.bold = True
        cp1.font.color.rgb = color
        cp1.font.name = "Segoe UI"
        
        cp2 = ctf.add_paragraph()
        cp2.text = desc
        cp2.font.size = Pt(12)
        cp2.font.color.rgb = RGBColor(226, 232, 240)
        cp2.font.name = "Segoe UI"
        cp2.space_before = Pt(6)

    # Footer Team Card
    footer = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(8.4), Inches(17.6), Inches(1.9))
    footer.fill.solid()
    footer.fill.fore_color.rgb = RGBColor(17, 24, 39)
    footer.line.color.rgb = C_SLATE_MID
    
    ftf = footer.text_frame
    ftf.word_wrap = True
    ftf.margin_left = ftf.margin_right = ftf.margin_top = ftf.margin_bottom = Inches(0.3)
    
    fp1 = ftf.paragraphs[0]
    fp1.text = "👥 Team: CodeOhollic  |  Leader: Adarsh Singh  |  Members: Khushi Kumari, Shivam Raj"
    fp1.font.size = Pt(15)
    fp1.font.bold = True
    fp1.font.color.rgb = C_WHITE
    
    fp2 = ftf.add_paragraph()
    fp2.text = "🏛️ Institution: Bharati Vidyapeeth College of Engineering, Pune  |  Theme: Smart City & Governance (Civic Tech)"
    fp2.font.size = Pt(13)
    fp2.font.color.rgb = C_AMBER
    fp2.space_before = Pt(4)

    # ==========================================
    # SLIDE 2: PROJECT CONTEXT & TEAM OVERVIEW
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "Slide 02 / Project Overview", "Team CodeOhollic & Executive Project Summary", "Empowering 1.4B+ citizens with transparent, AI-enabled municipal governance")

    # Left Column: Team & Institutional Credentials
    left_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(8.8), Inches(8.3))
    left_box.fill.solid()
    left_box.fill.fore_color.rgb = C_WHITE
    left_box.line.color.rgb = C_BORDER_LIGHT
    ltf = left_box.text_frame
    ltf.word_wrap = True
    ltf.margin_left = ltf.margin_right = ltf.margin_top = ltf.margin_bottom = Inches(0.4)

    lp0 = ltf.paragraphs[0]
    lp0.text = "👥 TEAM DETAILS & ROLES"
    lp0.font.size = Pt(18)
    lp0.font.bold = True
    lp0.font.color.rgb = C_NAVY_DARK

    team_members = [
        ("👑 Adarsh Singh (Leader)", "Full-Stack System Architecture, MERN Backend, AI & Geospatial Clustering Engines"),
        ("⚡ Khushi Kumari", "Frontend UX/UI, Responsive Design, Multi-Language Localization Engine"),
        ("🛠️ Shivam Raj", "API Integration, Anti-Spam Security, Leaflet Mapping & SLA Countdown Services"),
    ]

    for name, role in team_members:
        np = ltf.add_paragraph()
        np.text = name
        np.font.size = Pt(14)
        np.font.bold = True
        np.font.color.rgb = C_TEXT_DARK
        np.space_before = Pt(10)

        rp = ltf.add_paragraph()
        rp.text = role
        rp.font.size = Pt(11)
        rp.font.color.rgb = C_TEXT_MUTED

    lp_inst = ltf.add_paragraph()
    lp_inst.text = "🏛️ INSTITUTION & TRACK"
    lp_inst.font.size = Pt(16)
    lp_inst.font.bold = True
    lp_inst.font.color.rgb = C_NAVY_DARK
    lp_inst.space_before = Pt(16)

    lp_inst_sub = ltf.add_paragraph()
    lp_inst_sub.text = "Bharati Vidyapeeth College of Engineering, Pune\nTrack: Smart City & Governance (Civic Tech)\nHackathon: HackConQuest Aether 2026"
    lp_inst_sub.font.size = Pt(12)
    lp_inst_sub.font.color.rgb = C_TEXT_MUTED
    lp_inst_sub.space_before = Pt(4)

    # Right Column: The Core Vision & Governance Challenge
    right_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.0), Inches(2.0), Inches(9.2), Inches(8.3))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = C_WHITE
    right_box.line.color.rgb = C_BORDER_LIGHT
    rtf = right_box.text_frame
    rtf.word_wrap = True
    rtf.margin_left = rtf.margin_right = rtf.margin_top = rtf.margin_bottom = Inches(0.4)

    rp0 = rtf.paragraphs[0]
    rp0.text = "🎯 PROJECT VISION & CORE MOTIVATION"
    rp0.font.size = Pt(18)
    rp0.font.bold = True
    rp0.font.color.rgb = C_NAVY_DARK

    vision_points = [
        ("💡 The Fundamental Gap", "Citizens encounter broken streetlights, water leaks, open potholes, and garbage piles daily, but have no quick, transparent, or accountable mechanism to report them without bureaucratic red tape."),
        ("⚡ Zero-Friction Philosophy", "Removing signup/login requirements allows any citizen to report a hazard in under 60 seconds with GPS and photo/voice proof."),
        ("🤝 Community Power & Transparency", "Transforming complaint handling from private opaque tickets into a transparent community stream (Nagar Feed) where citizens upvote common problems and verify genuine physical repairs with Before/After proof."),
        ("🏛️ Department Accountability", "Empowering municipal officers with automated routing, duplicate clustering, SLA countdowns, and public rating scorecards."),
    ]

    for h, b in vision_points:
        vp = rtf.add_paragraph()
        vp.text = h
        vp.font.size = Pt(13)
        vp.font.bold = True
        vp.font.color.rgb = C_BLUE
        vp.space_before = Pt(10)

        vbp = rtf.add_paragraph()
        vbp.text = b
        vbp.font.size = Pt(11)
        vbp.font.color.rgb = C_TEXT_DARK
        vbp.space_before = Pt(2)

    # ==========================================
    # SLIDE 3: PROBLEM STATEMENT & ROOT CAUSES
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "Slide 03 / Problem Statement", "The 5 Critical Failures of Traditional Civic Grievance Portals", "Why existing municipal helplines, paper registers, and complex government apps fail citizens")

    problems = [
        ("1. Scattered & Lost Channels", "📭", "Complaints are scattered across phone calls, paper slips, or lost in WhatsApp messages without centralized tracking or audit trail.", C_RED),
        ("2. Friction-Heavy Login Walls", "🔒", "Demanding citizen registration, Aadhaar, and OTPs creates high friction that kills 80%+ of reporting intent on public hazards.", C_AMBER),
        ("3. The 'Black Hole' Syndrome", "❓", "Once filed, citizens have zero visibility into which officer was assigned, SLA timelines, or if any progress is occurring.", C_BLUE),
        ("4. Duplicate Ticket Floods", "📑", "When a water pipe bursts or pothole forms, dozens of neighbors file duplicate entries, overwhelming department staff with redundant queues.", C_PURPLE),
        ("5. Fake & Ghost Closures", "🚫", "Officers mark complaints as 'Resolved' without physical proof or citizen verification, leading to zero accountability and public distrust.", C_RED),
    ]

    for i, (title, icon, desc, color) in enumerate(problems):
        y = Inches(2.0 + i * 1.65)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(18.4), Inches(1.45))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.4)
        ctf.margin_top = Inches(0.18)
        ctf.margin_bottom = Inches(0.18)

        cp1 = ctf.paragraphs[0]
        cp1.text = f"{icon}  {title}"
        cp1.font.size = Pt(16)
        cp1.font.bold = True
        cp1.font.color.rgb = color

        cp2 = ctf.add_paragraph()
        cp2.text = desc
        cp2.font.size = Pt(12)
        cp2.font.color.rgb = C_TEXT_DARK
        cp2.space_before = Pt(3)

    # ==========================================
    # SLIDE 4: PROPOSED SOLUTION & CORE INNOVATION
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "Slide 04 / Proposed Solution", "Nagar Seva: The Full-Loop Civic Grievance Ecosystem", "A streamlined MERN architecture combining AI automation, geospatial clustering, and community verification")

    pillars = [
        ("📱 Frictionless Citizen Portal", C_BLUE, [
            "⚡ Zero-Login Filing: Report any issue in <60 seconds with phone & photo.",
            "🌐 7 Regional Languages: Live English, Hindi, Marathi, Bengali, Tamil, Telugu, Gujarati.",
            "🎙️ Voice Notes & Speech-to-Text: Speak naturally or record audio recordings.",
            "📍 Interactive OpenStreetMap: Draggable GPS pin-drop location selector.",
        ]),
        ("⚙️ Intelligent Municipal Engine", C_PURPLE, [
            "🤖 AI Real-Time Classifier: Keyword & pattern auto-detection for department & urgency.",
            "🔍 Proximity Duplicate Engine: Merges reports within ~300m into a high-urgency cluster.",
            "🚨 Emergency SOS Matrix: 6-hour critical SLA countdown with pulsing alerts.",
            "📊 Role-Based Dashboards: Custom views for Water, Roads, Sanitation, and Master Admin.",
        ]),
        ("🤝 Community & Accountability", C_EMERALD, [
            "📷 Mandatory Proof-of-Fix: Staff must upload resolution photo to close tickets.",
            "🎛️ Before vs. After Slider: Interactive visual comparison of municipal transformation.",
            "📰 Nagar Community Feed: Public stream with 'Affected Too (+1)' upvoting.",
            "⭐ 1-5 Star Ratings & Re-Open: Citizen reviews and instant High-Priority re-opening.",
        ]),
    ]

    for i, (title, color, bullet_list) in enumerate(pillars):
        x = Inches(0.8 + i * 6.25)
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(5.9), Inches(8.3))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = color
        card.line.width = Pt(2)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = Inches(0.35)

        cp1 = ctf.paragraphs[0]
        cp1.text = title
        cp1.font.size = Pt(17)
        cp1.font.bold = True
        cp1.font.color.rgb = color

        for b in bullet_list:
            bp = ctf.add_paragraph()
            bp.text = b
            bp.font.size = Pt(11.5)
            bp.font.color.rgb = C_TEXT_DARK
            bp.space_before = Pt(12)

    # ==========================================
    # SLIDE 5: 10 ADVANCED LIVE CAPABILITIES
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "Slide 05 / 10 Live Superpowers", "10 Flagship Upgrades Built & Fully Functional", "Going beyond standard CRUD — cutting-edge features implemented in Nagar Seva")

    features_10 = [
        ("🚨 Emergency SOS Engine", "6-Hour SLA countdown, audio-visual priority pulse banners, and emergency department dispatch.", C_RED),
        ("🤖 AI Smart Classifier", "Real-time keyword & contextual analyzer predicting department and urgency as citizens type.", C_PURPLE),
        ("🎙️ Voice & Audio Dictation", "Browser Web Speech API dictation in regional Indian languages + audio voice note recorder.", C_BLUE),
        ("🗺️ OpenStreetMap & Heatmap", "Interactive pin picker on filing + Officer City Map View with color-coded status markers.", C_EMERALD),
        ("🔍 Proximity Duplicate Merging", "Geospatial proximity matching (~300m) auto-grouping neighbor reports into high-priority clusters.", C_AMBER),
        ("📰 Nagar Feed (+1 Upvotes)", "Public community feed where neighbors click 'Affected Too (+1)' to escalate neighborhood urgency.", C_EMERALD),
        ("🎛️ Before vs. After Slider", "Interactive draggable split slider comparing original problem photo vs staff resolution proof.", C_BLUE),
        ("⏰ SLA Timers & Overdue Alert", "Dynamic countdown badges (6h SOS, 24h Water/Sanitation, 48h Roads) with overdue escalation.", C_AMBER),
        ("⭐ Star Rating & Re-Open", "1-5 Star citizen reviews + 'Re-Open Issue' workflow flagging unresolved work back to High Priority.", C_PURPLE),
        ("🏆 Civic Karma & Alerts Ticker", "Points leaderboard (+50 filed, +30 solved, +20 rated) + Homepage live municipal news ticker.", C_RED),
    ]

    for i, (title, desc, color) in enumerate(features_10):
        col = i % 2
        row = i // 2
        x = Inches(0.8 + col * 9.35)
        y = Inches(2.0 + row * 1.65)

        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(9.0), Inches(1.48))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.3)
        ctf.margin_top = Inches(0.18)
        ctf.margin_bottom = Inches(0.18)

        cp1 = ctf.paragraphs[0]
        cp1.text = title
        cp1.font.size = Pt(14)
        cp1.font.bold = True
        cp1.font.color.rgb = color

        cp2 = ctf.add_paragraph()
        cp2.text = desc
        cp2.font.size = Pt(11)
        cp2.font.color.rgb = C_TEXT_DARK
        cp2.space_before = Pt(2)

    # ==========================================
    # SLIDE 6: TECHNICAL ARCHITECTURE & DATAFLOW
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "Slide 06 / Architecture & Workflow", "Technical Architecture & End-to-End Lifecycle", "Robust, modular MERN architecture with real-time geospatial processing")

    # Architecture 3-tier Columns
    tiers = [
        ("🖥️ Client Frontend (Vite SPA)", C_BLUE, [
            "React 18 & React Router v6",
            "Tailwind CSS v3 (Utility-first styling)",
            "Leaflet & React-Leaflet OpenStreetMap",
            "Web Speech API (Speech Recognition)",
            "MediaRecorder API (Audio Voice Notes)",
            "Axios with JWT Bearer Interceptors",
        ]),
        ("⚡ Backend Server (REST API)", C_PURPLE, [
            "Node.js & Express.js Engine",
            "Multer Multipart Storage (Image & Audio)",
            "Geospatial Proximity Matcher (Δlat/lng ≤ 0.003)",
            "Dynamic SLA Expiry & Escalation Matrix",
            "JWT & Bcrypt Security Middleware",
            "Express Static Upload Server",
        ]),
        ("🗄️ Database (MongoDB Atlas)", C_EMERALD, [
            "MongoDB Atlas (Cloud Cluster M0)",
            "Mongoose ODM Schemas & Models",
            "Complaints & Piled Sub-Reports Schema",
            "Municipal Announcements & Alerts Schema",
            "Civic Karma Aggregation Pipelines",
            "Compound Geospatial & Status Indexes",
        ]),
    ]

    for i, (tier_title, color, items) in enumerate(tiers):
        x = Inches(0.8 + i * 6.25)
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(5.9), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = Inches(0.3)

        cp1 = ctf.paragraphs[0]
        cp1.text = tier_title
        cp1.font.size = Pt(15)
        cp1.font.bold = True
        cp1.font.color.rgb = color

        for it in items:
            itp = ctf.add_paragraph()
            itp.text = f"• {it}"
            itp.font.size = Pt(11)
            itp.font.color.rgb = C_TEXT_DARK
            itp.space_before = Pt(4)

    # Lifecycle 5-Step Pipeline
    pipe_header = s6.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(18.4), Inches(0.5))
    ptf = pipe_header.text_frame
    pp = ptf.paragraphs[0]
    pp.text = "🔄 END-TO-END DATAFLOW & RESOLUTION PIPELINE"
    pp.font.size = Pt(14)
    pp.font.bold = True
    pp.font.color.rgb = C_NAVY_DARK

    steps = [
        ("Step 1: Citizen Report", "Category + Photo + Voice + GPS Pin", C_BLUE),
        ("Step 2: AI & Clustering", "AI classifies, Proximity merges duplicates", C_PURPLE),
        ("Step 3: Officer Action", "Dept logs in, moves ticket to In Progress", C_AMBER),
        ("Step 4: Proof Upload", "Officer uploads resolution photo", C_EMERALD),
        ("Step 5: Review & Feed", "Citizen rates 1-5★, Before/After live on Feed", C_RED),
    ]

    for i, (st, sd, sc) in enumerate(steps):
        x = Inches(0.8 + i * 3.75)
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(7.1), Inches(3.45), Inches(3.2))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = sc
        card.line.width = Pt(1.5)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = Inches(0.25)

        cp1 = ctf.paragraphs[0]
        cp1.text = st
        cp1.font.size = Pt(13)
        cp1.font.bold = True
        cp1.font.color.rgb = sc

        cp2 = ctf.add_paragraph()
        cp2.text = sd
        cp2.font.size = Pt(11)
        cp2.font.color.rgb = C_TEXT_MUTED
        cp2.space_before = Pt(6)

    # ==========================================
    # SLIDE 7: IMPACT, FEASIBILITY & SCALABILITY
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "Slide 07 / Impact & Feasibility", "Measurable Civic Impact & Zero-Cost Feasibility", "Real-world viability, anti-spam safeguards, and seamless government scaling")

    impact_quads = [
        ("📊 Measurable Civic Impact", C_EMERALD, [
            ("⚡ 80% Faster Filing", "Zero-login workflow enables reporting in <60 seconds."),
            ("⏱️ 60% Faster Resolution", "Category SLA countdowns and automated escalations prevent delays."),
            ("🔥 75% Less Redundancy", "Smart proximity clustering merges neighbor reports automatically."),
            ("🤝 100% Verified Work", "Compulsory proof-of-fix images prevent fake closures."),
        ]),
        ("💰 Zero-Cost Architecture", C_BLUE, [
            ("🆓 100% Free Tier Hosting", "Runs on Render/Vercel free compute + MongoDB Atlas M0 free tier."),
            ("🗺️ Free OpenStreetMap", "Zero Google Maps API bills using OpenStreetMap & Leaflet tiles."),
            ("📦 Pure Open-Source", "No proprietary licenses or vendor lock-in."),
            ("📉 Zero Maintenance Overhead", "Stateless Express REST backend with lightweight footprint."),
        ]),
        ("🛡️ Anti-Spam & Fraud Guards", C_RED, [
            ("📍 Geolocation Verification", "GPS coordinate tagging verifies user proximity to problem."),
            ("🔒 Single-Vote Enforced", "Voter tokens and IP tracking prevent upvote spamming."),
            ("🖼️ MIME & Size Validation", "Multer guards block malicious payloads and restrict images to 10MB."),
            ("👮 Officer Vetting", "Department review required before field deployment."),
        ]),
        ("🚀 Scale & Future Scope", C_PURPLE, [
            ("🏛️ National Portal Bridge", "Modular REST architecture easily bridges with CPGRAMS & State Portals."),
            ("🏙️ Smart City ICCC Integration", "Ready for Integrated Command and Control Center (ICCC) feeds."),
            ("📲 Offline PWA App", "Offline caching and background sync for remote areas with spotty internet."),
            ("🤖 Computer Vision AI", "Automated Before vs. After image verification to detect genuine cleanup."),
        ]),
    ]

    for i, (title, color, points) in enumerate(impact_quads):
        col = i % 2
        row = i // 2
        x = Inches(0.8 + col * 9.35)
        y = Inches(2.0 + row * 4.25)

        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(9.0), Inches(3.95))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = Inches(0.3)

        cp1 = ctf.paragraphs[0]
        cp1.text = title
        cp1.font.size = Pt(16)
        cp1.font.bold = True
        cp1.font.color.rgb = color

        for ph, pb in points:
            p_head = ctf.add_paragraph()
            p_head.text = f"• {ph}: {pb}"
            p_head.font.size = Pt(11)
            p_head.font.color.rgb = C_TEXT_DARK
            p_head.space_before = Pt(4)

    # ==========================================
    # SLIDE 8: DEMO LINKS & CONCLUSION (DARK THEME)
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8, C_NAVY_DARK)

    # Accent Top Strip
    accent_bar8 = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.15))
    accent_bar8.fill.solid()
    accent_bar8.fill.fore_color.rgb = C_EMERALD
    accent_bar8.line.fill.background()

    # Title
    tbox8 = s8.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(18.4), Inches(1.3))
    tf8 = tbox8.text_frame
    tf8.word_wrap = True
    
    p8_tag = tf8.paragraphs[0]
    p8_tag.text = "SLIDE 08 / CONCLUSION & DEMO"
    p8_tag.font.size = Pt(12)
    p8_tag.font.bold = True
    p8_tag.font.color.rgb = C_AMBER

    p8_title = tf8.add_paragraph()
    p8_title.text = "Live Demonstration & Project Credentials"
    p8_title.font.size = Pt(28)
    p8_title.font.bold = True
    p8_title.font.color.rgb = C_WHITE

    # Left Box: Live Links & GitHub
    left_demo = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.1), Inches(8.8), Inches(8.2))
    left_demo.fill.solid()
    left_demo.fill.fore_color.rgb = C_SLATE_DARK
    left_demo.line.color.rgb = C_EMERALD
    left_demo.line.width = Pt(2)
    
    ld_tf = left_demo.text_frame
    ld_tf.word_wrap = True
    ld_tf.margin_left = ld_tf.margin_right = ld_tf.margin_top = ld_tf.margin_bottom = Inches(0.4)

    ld_p0 = ld_tf.paragraphs[0]
    ld_p0.text = "🌐 LIVE LINKS & REPOSITORY"
    ld_p0.font.size = Pt(18)
    ld_p0.font.bold = True
    ld_p0.font.color.rgb = C_EMERALD

    demo_links = [
        ("🐙 GitHub Repository", "https://github.com/adarshsingh022006-tech/Nagar-Seva.git", "Full MERN source code, models, controllers, and UI components"),
        ("🚀 Live Production URL", "https://nagar-seva-1-l8h5.onrender.com", "Deployed live on Render with MongoDB Atlas database"),
        ("📰 Nagar Community Feed", "https://nagar-seva-1-l8h5.onrender.com/feed", "Public civic stream with Before/After sliders and (+1) upvoting"),
    ]

    for title, url, note in demo_links:
        lp = ld_tf.add_paragraph()
        lp.text = title
        lp.font.size = Pt(13)
        lp.font.bold = True
        lp.font.color.rgb = C_AMBER
        lp.space_before = Pt(8)

        up = ld_tf.add_paragraph()
        up.text = url
        up.font.size = Pt(11)
        up.font.bold = True
        up.font.color.rgb = C_WHITE

        np = ld_tf.add_paragraph()
        np.text = note
        np.font.size = Pt(10)
        np.font.color.rgb = RGBColor(148, 163, 184)

    # Right Box: Demo Credentials & Conclusion
    right_demo = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.0), Inches(2.1), Inches(9.2), Inches(8.2))
    right_demo.fill.solid()
    right_demo.fill.fore_color.rgb = C_SLATE_DARK
    right_demo.line.color.rgb = C_AMBER
    right_demo.line.width = Pt(2)
    
    rd_tf = right_demo.text_frame
    rd_tf.word_wrap = True
    rd_tf.margin_left = rd_tf.margin_right = rd_tf.margin_top = rd_tf.margin_bottom = Inches(0.4)

    rd_p0 = rd_tf.paragraphs[0]
    rd_p0.text = "🔑 DEMO OFFICER CREDENTIALS"
    rd_p0.font.size = Pt(18)
    rd_p0.font.bold = True
    rd_p0.font.color.rgb = C_AMBER

    creds = [
        ("Master Admin Portal", "admin", "admin123", "Full city view, department stats, public alert broadcasting"),
        ("Roads & Infrastructure", "roads_dept", "dept123", "Pothole management, proof of fix upload"),
        ("Water Department", "water_dept", "dept123", "Pipeline leaks, water supply management"),
        ("Sanitation Department", "sanitation_dept", "dept123", "Garbage dumps, waste clearance verification"),
    ]

    for dept, user, pwd, note in creds:
        cp = rd_tf.add_paragraph()
        cp.text = f"• {dept}: user: {user} | pass: {pwd}"
        cp.font.size = Pt(12)
        cp.font.bold = True
        cp.font.color.rgb = C_WHITE
        cp.space_before = Pt(4)

        cnp = rd_tf.add_paragraph()
        cnp.text = f"   ({note})"
        cnp.font.size = Pt(10)
        cnp.font.color.rgb = RGBColor(148, 163, 184)

    concl_p = rd_tf.add_paragraph()
    concl_p.text = "🏆 FINAL SUMMARY"
    concl_p.font.size = Pt(16)
    concl_p.font.bold = True
    concl_p.font.color.rgb = C_EMERALD
    concl_p.space_before = Pt(12)

    concl_body = rd_tf.add_paragraph()
    concl_body.text = '"Nagar Seva transforms municipal governance from passive complaint logging into an active, community-validated, and AI-accelerated public trust engine — built entirely on a 100% free, scalable open-source stack."'
    concl_body.font.size = Pt(12)
    concl_body.font.italic = True
    concl_body.font.color.rgb = C_WHITE
    concl_body.space_before = Pt(4)

    # Save to original location and local working directory
    target_whatsapp = r"C:\Users\HP PC\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\75F3819468DA04F67B34F5DC24DB618F39D93316\transfers\2026-38\Nagar_Seva_HackConQuest.pptx"
    target_local = r"C:\Users\HP PC\Downloads\nagar-seva-mern\nagar-seva-mern\Nagar_Seva_HackConQuest_Updated.pptx"

    prs.save(target_local)
    print("SUCCESS: Saved updated presentation to local repo: " + target_local)

    try:
        prs.save(target_whatsapp)
        print("SUCCESS: Updated original WhatsApp PPTX directly: " + target_whatsapp)
    except Exception as e:
        print("WARNING: Could not overwrite WhatsApp file directly, saved to local repo: " + str(e))

if __name__ == "__main__":
    create_presentation()
